"""Read and validate the issue #13 staging workbook before anything is stored.

The workbook is a human/AI research staging package: 115 bibliography records,
candidate scientific claims, the scientific field contract and team research-gap
hypotheses. None of it is independently verified. Validation is therefore about
*integrity* (nothing lost, nothing dangling, nothing silently promoted), never
about scientific correctness - that belongs to the #19 evidence gate and the
independent Scientific Auditor.

Every check that fails is a hard error and the import writes nothing. A
workbook that disagrees with itself in a way that must stay visible (two
sheets reporting different audit states for one paper) is imported with a
warning that preserves both values; no side is silently chosen.
"""

from __future__ import annotations

import ast
import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

REQUIRED_SHEETS: tuple[str, ...] = (
    "START_HERE", "PAPER_INDEX", "SCIENTIFIC_CLAIMS", "RESEARCH_GAPS",
    "PAPER_RECORDS", "FIELD_CONTRACT", "README", "AUDIT_PROTOCOL", "METHODS_GUIDE",
)
DOCUMENT_SHEETS: tuple[str, ...] = ("README", "AUDIT_PROTOCOL", "METHODS_GUIDE")

# A supplement workbook (for example Ocean_Research_Intelligence_EXPANDED_V2)
# adds candidate claims to a corpus that a base workbook already imported. It
# carries no PAPER_RECORDS and no FIELD_CONTRACT, so it cannot introduce a
# paper and is validated against the stored base corpus at import time.
BASE = "base"
SUPPLEMENT = "supplement"
SUPPLEMENT_SHEETS: tuple[str, ...] = (
    "START_HERE", "PAPER_INDEX", "SCIENTIFIC_CLAIMS", "RESEARCH_GAPS", "COVERAGE",
)

REQUIRED_COLUMNS: dict[str, tuple[str, ...]] = {
    "PAPER_INDEX": (
        "paper_id", "title", "year", "doi", "domain", "source_url", "review_stage",
        "source_edition", "scientific_extraction_status", "scientific_audit_status",
    ),
    "SCIENTIFIC_CLAIMS": (
        "claim_id", "paper_id", "experiment_id", "field_path", "value", "subject_scope",
        "claim_type", "evidence_source_url", "source_doi", "source_edition",
        "source_section", "source_locator", "verbatim_evidence", "pdf_page",
        "evidence_review", "scientific_status", "independent_audit", "notes",
    ),
    "RESEARCH_GAPS": (
        "candidate_topic", "testable_question", "basis_paper_ids", "status", "qualification",
    ),
    "PAPER_RECORDS": ("paper_id", "record_json"),
    "FIELD_CONTRACT": ("group", "field_name", "field_path"),
    "COVERAGE": ("paper_id", "claim_count", "coverage_level", "audit_state"),
    "NEW_DETAILED_PAPERS": ("paper_id", "claims_added"),
}

# A staging package may carry candidate states only. VERIFIED needs an
# independent attestation this package does not contain, and NOT_REPORTED
# needs a documented complete search scope, which it does not contain either
# ("Never infer NOT_REPORTED from an abstract", AUDIT_PROTOCOL step 6).
ALLOWED_CLAIM_STATUSES = frozenset({"NOT_VERIFIED", "EXTRACTION_ERROR", "CONFLICT"})
REFUSED_CLAIM_STATUSES = {
    "VERIFIED": "a staging package cannot carry an independent attestation",
    "PARTIALLY_VERIFIED": "a staging package cannot carry an independent attestation",
    "NOT_REPORTED": "absence requires a documented complete search scope, which staging rows do not carry",
}
# Provenance classes defined in AGENTS.md. Other values are preserved verbatim
# and reported, never remapped.
PROJECT_PROVENANCE = frozenset({
    "AUTHOR_REPORTED_FACT", "AUTHOR_REPORTED_LIMITATION", "AI_INTERPRETATION", "TEAM_NOTE",
})
# START_HERE declares the package's own counts; a mismatch means a truncated
# or edited sheet and the import is refused.
DECLARED_COUNTS = {
    "Corpus records": "PAPER_INDEX",
    "Raw paper records": "PAPER_RECORDS",
    "Scientific claim candidates": "SCIENTIFIC_CLAIMS",
    "Scientific field definitions": "FIELD_CONTRACT",
    "Research-gap candidates": "RESEARCH_GAPS",
}
SUPPLEMENT_DECLARED_COUNTS = {
    "Papers indexed": "PAPER_INDEX",
    "Claims total": "SCIENTIFIC_CLAIMS",
    "Detailed papers now": "DETAILED_PAPERS",
    "New detailed papers": "NEW_DETAILED_PAPERS",
}


class WorkbookValidationError(ValueError):
    """The workbook failed integrity validation; nothing may be imported."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__(f"{len(errors)} workbook integrity error(s): " + "; ".join(errors[:10]))
        self.errors = errors


@dataclass(frozen=True)
class StagingWorkbook:
    source_name: str
    sha256: str
    sheet_names: list[str]
    start_here: dict[str, str]
    paper_index: list[dict[str, Any]]
    claims: list[dict[str, Any]]
    research_gaps: list[dict[str, Any]]
    paper_records: list[dict[str, Any]]
    field_contract: list[dict[str, Any]]
    documents: dict[str, str]
    warnings: list[str] = field(default_factory=list)
    kind: str = BASE
    coverage: list[dict[str, Any]] = field(default_factory=list)
    new_detailed_papers: list[dict[str, Any]] = field(default_factory=list)


def normalize_doi(doi: str | None) -> str | None:
    if not doi:
        return None
    value = doi.strip().lower()
    value = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)", "", value)
    return value or None


def normalize_title(title: str | None) -> str | None:
    """Case/punctuation-insensitive title key. Never used to merge different DOIs."""
    if not title:
        return None
    text = unicodedata.normalize("NFKD", title)
    text = "".join(character for character in text if not unicodedata.combining(character))
    key = re.sub(r"[\W_]+", " ", text.casefold()).strip()
    return key or None


def parse_claim_value(raw: Any) -> Any:
    """Return a list for a literal list of strings, else the raw value unchanged.

    The workbook stores list-valued claims as Python list literals
    ("['nadir altimetry', 'wide-swath altimetry']"). ``ast.literal_eval`` only
    evaluates literals, so this cannot execute code; anything that is not a
    literal list of strings is kept exactly as written.
    """
    if not isinstance(raw, str):
        return raw
    text = raw.strip()
    if not (text.startswith("[") and text.endswith("]")):
        return raw
    try:
        parsed = ast.literal_eval(text)
    except (ValueError, SyntaxError):
        return raw
    if isinstance(parsed, list) and all(isinstance(item, str) for item in parsed):
        return parsed
    return raw


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


def read_workbook(path: str | Path) -> StagingWorkbook:
    """Load every sheet and validate the package's integrity, or raise."""
    from openpyxl import load_workbook

    source = Path(path)
    content = source.read_bytes()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        errors: list[str] = [
            f"missing worksheet {name}" for name in REQUIRED_SHEETS if name not in sheet_names
        ]
        kind = BASE
        if errors and all(name in sheet_names for name in SUPPLEMENT_SHEETS) and not (
            {"PAPER_RECORDS", "FIELD_CONTRACT"} & set(sheet_names)
        ):
            kind, errors = SUPPLEMENT, []
        if errors:
            raise WorkbookValidationError(errors)

        def rows(name: str) -> list[tuple[Any, ...]]:
            return [
                tuple(_cell(value) for value in row)
                for row in workbook[name].iter_rows(values_only=True)
                if any(value is not None and str(value).strip() for value in row)
            ]

        def table(name: str) -> list[dict[str, Any]]:
            data = rows(name)
            if not data:
                errors.append(f"worksheet {name} is empty")
                return []
            header = [str(value) if value is not None else "" for value in data[0]]
            missing = [column for column in REQUIRED_COLUMNS.get(name, ()) if column not in header]
            if missing:
                errors.append(f"worksheet {name} lacks column(s) {missing}")
            return [
                {column: row[index] if index < len(row) else None for index, column in enumerate(header) if column}
                for row in data[1:]
            ]

        start_here = {
            str(row[0]): "" if len(row) < 2 or row[1] is None else str(row[1])
            for row in rows("START_HERE")[1:] if row and row[0] is not None
        }
        documents = {
            name: "\n".join(
                "" if len(row) < 2 or row[1] is None else str(row[1]) for row in rows(name)[1:]
            )
            for name in DOCUMENT_SHEETS if kind == BASE
        }
        paper_index = table("PAPER_INDEX")
        claims = table("SCIENTIFIC_CLAIMS")
        gaps = table("RESEARCH_GAPS")
        record_rows = table("PAPER_RECORDS") if kind == BASE else []
        contract = table("FIELD_CONTRACT") if kind == BASE else []
        coverage = table("COVERAGE") if kind == SUPPLEMENT else []
        new_detailed = (
            table("NEW_DETAILED_PAPERS")
            if kind == SUPPLEMENT and "NEW_DETAILED_PAPERS" in sheet_names else []
        )
    finally:
        workbook.close()

    records: list[dict[str, Any]] = []
    for position, row in enumerate(record_rows, start=2):
        try:
            record = json.loads(row.get("record_json") or "")
        except json.JSONDecodeError as exc:
            errors.append(f"PAPER_RECORDS row {position}: record_json is not valid JSON ({exc.msg})")
            continue
        if not isinstance(record, dict):
            errors.append(f"PAPER_RECORDS row {position}: record_json is not an object")
            continue
        if record.get("paper_id") != row.get("paper_id"):
            errors.append(
                f"PAPER_RECORDS row {position}: paper_id column {row.get('paper_id')!r} "
                f"differs from record_json paper_id {record.get('paper_id')!r}"
            )
        records.append(record)

    staged = StagingWorkbook(
        source_name=source.name,
        sha256=hashlib.sha256(content).hexdigest(),
        sheet_names=sheet_names,
        start_here=start_here,
        paper_index=paper_index,
        claims=claims,
        research_gaps=gaps,
        paper_records=records,
        field_contract=contract,
        documents=documents,
        kind=kind,
        coverage=coverage,
        new_detailed_papers=new_detailed,
    )
    if kind == SUPPLEMENT:
        errors.extend(_supplement_errors(staged))
    else:
        errors.extend(_integrity_errors(staged))
    if errors:
        raise WorkbookValidationError(errors)
    staged.warnings.extend(_integrity_warnings(staged))
    return staged


def _duplicates(values: list[Any]) -> list[Any]:
    seen: set[Any] = set()
    repeated: list[Any] = []
    for value in values:
        if value in seen and value not in repeated:
            repeated.append(value)
        seen.add(value)
    return repeated


def _integrity_errors(book: StagingWorkbook) -> list[str]:
    errors: list[str] = []
    counts = {
        "PAPER_INDEX": len(book.paper_index), "PAPER_RECORDS": len(book.paper_records),
        "SCIENTIFIC_CLAIMS": len(book.claims), "FIELD_CONTRACT": len(book.field_contract),
        "RESEARCH_GAPS": len(book.research_gaps),
    }
    for label, sheet in DECLARED_COUNTS.items():
        declared = book.start_here.get(label)
        if declared is not None and declared.strip().isdigit() and int(declared) != counts[sheet]:
            errors.append(f"START_HERE declares {label} = {declared} but {sheet} has {counts[sheet]} rows")

    # Field contract.
    paths = [row.get("field_path") for row in book.field_contract]
    if any(not path for path in paths):
        errors.append("FIELD_CONTRACT has a row without field_path")
    for path in _duplicates(paths):
        errors.append(f"FIELD_CONTRACT defines {path!r} more than once")
    for row in book.field_contract:
        if row.get("field_path") and row.get("field_path") != f"{row.get('group')}.{row.get('field_name')}":
            errors.append(
                f"FIELD_CONTRACT {row.get('field_path')!r} does not equal group.field_name "
                f"({row.get('group')}.{row.get('field_name')})"
            )
    contract = set(paths)

    # Papers: identity and cross-sheet agreement.
    errors.extend(_paper_identity_errors(book))
    index_ids = [row.get("paper_id") for row in book.paper_index]
    record_ids = [record.get("paper_id") for record in book.paper_records]
    for paper_id in _duplicates(record_ids):
        errors.append(f"PAPER_RECORDS lists paper_id {paper_id!r} more than once")
    if set(index_ids) != set(record_ids):
        errors.append(
            "PAPER_INDEX and PAPER_RECORDS disagree on the paper set: only in index "
            f"{sorted(set(index_ids) - set(record_ids))}, only in records "
            f"{sorted(set(record_ids) - set(index_ids))}"
        )
    records = {record.get("paper_id"): record for record in book.paper_records}
    for row in book.paper_index:
        record = records.get(row.get("paper_id"))
        if record is None:
            continue
        for column in ("title", "doi", "year"):
            index_value = row.get(column)
            record_value = _cell(record.get(column))
            if (str(index_value) if index_value is not None else None) != (
                str(record_value) if record_value is not None else None
            ):
                errors.append(
                    f"{row.get('paper_id')}: PAPER_INDEX {column} {index_value!r} differs from "
                    f"PAPER_RECORDS {column} {record_value!r}"
                )

    # Records must carry exactly the contract's fields.
    slot_of_claim: dict[str, tuple[str, str]] = {}
    for record in book.paper_records:
        paper_id = record.get("paper_id")
        slots = record.get("field_claim_ids")
        if not isinstance(slots, dict):
            errors.append(f"{paper_id}: record has no field_claim_ids object")
            continue
        record_paths: set[str] = set()
        for group, fields in slots.items():
            if not isinstance(fields, dict):
                errors.append(f"{paper_id}: field group {group!r} is not an object")
                continue
            for name, claim_ids in fields.items():
                path = f"{group}.{name}"
                record_paths.add(path)
                if not isinstance(claim_ids, list) or not all(isinstance(item, str) for item in claim_ids):
                    errors.append(f"{paper_id}: {path} claim ids are not a list of strings")
                    continue
                for claim_id in claim_ids:
                    if claim_id in slot_of_claim:
                        errors.append(f"claim {claim_id} is referenced by several record fields")
                    slot_of_claim[claim_id] = (paper_id, path)
        if record_paths != contract:
            errors.append(
                f"{paper_id}: record fields differ from FIELD_CONTRACT (missing "
                f"{sorted(contract - record_paths)[:5]}, extra {sorted(record_paths - contract)[:5]})"
            )

    # Claims: referential integrity and status discipline.
    errors.extend(_claim_row_errors(book))
    claim_ids = [row.get("claim_id") for row in book.claims]
    for row in book.claims:
        claim_id = row.get("claim_id")
        if not claim_id:
            continue
        if row.get("field_path") not in contract:
            errors.append(f"claim {claim_id} uses field_path {row.get('field_path')!r} outside FIELD_CONTRACT")
        slot = slot_of_claim.get(claim_id)
        if slot is None:
            errors.append(f"claim {claim_id} is not referenced by any PAPER_RECORDS field")
        elif slot != (row.get("paper_id"), row.get("field_path")):
            errors.append(
                f"claim {claim_id} is filed under {slot} in PAPER_RECORDS but declares "
                f"({row.get('paper_id')}, {row.get('field_path')})"
            )
    for claim_id in sorted(set(slot_of_claim) - set(claim_ids)):
        errors.append(f"PAPER_RECORDS references claim {claim_id} which SCIENTIFIC_CLAIMS does not contain")

    errors.extend(_gap_errors(book))
    return errors


def _paper_identity_errors(book: StagingWorkbook) -> list[str]:
    errors: list[str] = []
    index_ids = [row.get("paper_id") for row in book.paper_index]
    for paper_id in _duplicates(index_ids):
        errors.append(f"PAPER_INDEX lists paper_id {paper_id!r} more than once")
    if any(not paper_id for paper_id in index_ids):
        errors.append("PAPER_INDEX has a row without paper_id")
    doi_keys = [normalize_doi(row.get("doi")) for row in book.paper_index]
    for doi in _duplicates([key for key in doi_keys if key]):
        owners = [row.get("paper_id") for row in book.paper_index if normalize_doi(row.get("doi")) == doi]
        errors.append(f"DOI {doi} is assigned to several paper_ids {owners}; deduplicate before import")
    title_keys = [(normalize_title(row.get("title")), str(row.get("year"))) for row in book.paper_index]
    for key in _duplicates([key for key in title_keys if key[0]]):
        owners = [
            row.get("paper_id") for row in book.paper_index
            if (normalize_title(row.get("title")), str(row.get("year"))) == key
        ]
        errors.append(f"title/year {key} is assigned to several paper_ids {owners}; deduplicate before import")
    return errors


def _claim_row_errors(book: StagingWorkbook) -> list[str]:
    """Checks every claim row must pass whatever the workbook layout."""
    errors: list[str] = []
    for claim_id in _duplicates([row.get("claim_id") for row in book.claims]):
        errors.append(f"SCIENTIFIC_CLAIMS lists claim_id {claim_id!r} more than once")
    paper_set = {row.get("paper_id") for row in book.paper_index}
    for row in book.claims:
        claim_id = row.get("claim_id")
        if not claim_id:
            errors.append("SCIENTIFIC_CLAIMS has a row without claim_id")
            continue
        if row.get("paper_id") not in paper_set:
            errors.append(f"claim {claim_id} references unknown paper {row.get('paper_id')!r}")
        if not row.get("field_path"):
            errors.append(f"claim {claim_id} has no field_path")
        status = row.get("scientific_status")
        if status in REFUSED_CLAIM_STATUSES:
            errors.append(f"claim {claim_id} is {status}: {REFUSED_CLAIM_STATUSES[status]}")
        elif status not in ALLOWED_CLAIM_STATUSES:
            errors.append(f"claim {claim_id} has unknown scientific_status {status!r}")
        if row.get("value") is None:
            errors.append(f"claim {claim_id} has no value")
    return errors


def _supplement_errors(book: StagingWorkbook) -> list[str]:
    """Integrity of a supplement workbook on its own.

    Checks that need the stored base corpus (every paper already known, field
    paths inside the stored contract) run in ``CorpusRepository.import_workbook``.
    """
    errors = _paper_identity_errors(book) + _claim_row_errors(book) + _gap_errors(book)
    per_paper: dict[Any, int] = {}
    for row in book.claims:
        per_paper[row.get("paper_id")] = per_paper.get(row.get("paper_id"), 0) + 1
    counts = {
        "PAPER_INDEX": len(book.paper_index), "SCIENTIFIC_CLAIMS": len(book.claims),
        "DETAILED_PAPERS": len(per_paper), "NEW_DETAILED_PAPERS": len(book.new_detailed_papers),
    }
    for label, sheet in SUPPLEMENT_DECLARED_COUNTS.items():
        declared = book.start_here.get(label)
        if declared is not None and declared.strip().isdigit() and int(declared) != counts[sheet]:
            errors.append(f"START_HERE declares {label} = {declared} but the workbook has {counts[sheet]}")
    before, added, total = (book.start_here.get(label, "") for label in ("Claims before", "Claims added", "Claims total"))
    if all(value.strip().isdigit() for value in (before, added, total)) and int(before) + int(added) != int(total):
        errors.append(f"START_HERE declares Claims before {before} + Claims added {added} != Claims total {total}")

    index_ids = {row.get("paper_id") for row in book.paper_index}
    coverage_ids = [row.get("paper_id") for row in book.coverage]
    for paper_id in _duplicates(coverage_ids):
        errors.append(f"COVERAGE lists paper_id {paper_id!r} more than once")
    if set(coverage_ids) != index_ids:
        errors.append(
            "COVERAGE and PAPER_INDEX disagree on the paper set: only in COVERAGE "
            f"{sorted(map(str, set(coverage_ids) - index_ids))}, only in PAPER_INDEX "
            f"{sorted(map(str, index_ids - set(coverage_ids)))}"
        )
    for row in book.coverage:
        actual = per_paper.get(row.get("paper_id"), 0)
        if str(row.get("claim_count")) != str(actual):
            errors.append(
                f"COVERAGE declares {row.get('claim_count')} claims for {row.get('paper_id')} "
                f"but SCIENTIFIC_CLAIMS has {actual}"
            )
    for row in book.new_detailed_papers:
        if row.get("paper_id") not in index_ids:
            errors.append(f"NEW_DETAILED_PAPERS lists unknown paper {row.get('paper_id')!r}")
    return errors


def _gap_errors(book: StagingWorkbook) -> list[str]:
    errors: list[str] = []
    paper_set = {row.get("paper_id") for row in book.paper_index}
    for row in book.research_gaps:
        if not row.get("candidate_topic"):
            errors.append("RESEARCH_GAPS has a row without candidate_topic")
        basis = [item.strip() for item in str(row.get("basis_paper_ids") or "").split(";") if item.strip()]
        unknown = [item for item in basis if item not in paper_set]
        if unknown:
            errors.append(f"research gap {row.get('candidate_topic')!r} cites unknown papers {unknown}")
    for topic in _duplicates([normalize_title(row.get("candidate_topic")) for row in book.research_gaps]):
        errors.append(f"RESEARCH_GAPS lists the candidate topic {topic!r} more than once")
    return errors


def _integrity_warnings(book: StagingWorkbook) -> list[str]:
    warnings: list[str] = []
    records = {record["paper_id"]: record for record in book.paper_records}
    if book.kind == SUPPLEMENT:
        warnings.append(
            "supplement workbook: no PAPER_RECORDS, so stored records are kept as supplied and "
            "new claims are placed by field_path; COVERAGE coverage_level/audit_state are stored "
            "beside the PAPER_INDEX and PAPER_RECORDS states, never in place of them"
        )
    for row in book.paper_index:
        record = records.get(row["paper_id"])
        if record is None:
            continue
        if row.get("scientific_audit_status") != record.get("scientific_audit_status"):
            warnings.append(
                f"{row['paper_id']}: PAPER_INDEX scientific_audit_status "
                f"{row.get('scientific_audit_status')!r} differs from PAPER_RECORDS "
                f"{record.get('scientific_audit_status')!r}; both are preserved, neither is chosen"
            )
    other = sorted({
        str(row.get("claim_type")) for row in book.claims
        if row.get("claim_type") not in PROJECT_PROVENANCE
    })
    if other:
        warnings.append(
            f"claim_type value(s) outside the AGENTS.md provenance classes are preserved "
            f"verbatim, not remapped: {other}"
        )
    return warnings
