"""Issue #13: importing the staging corpus workbook.

The workbook is unaudited research staging material. These tests pin the
properties that make importing it safe: nothing is lost, nothing is promoted,
missing fields are NOT_EXTRACTED rather than NOT_REPORTED, no page or quotation
is invented, re-imports are idempotent, duplicates resolve to one work, and an
independently audited row - or the canonical PaperRecord table - is never
overwritten.
"""

from __future__ import annotations

import copy
import json
import sqlite3
from dataclasses import replace
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ocean_research_hub.api import create_app
from ocean_research_hub.corpus.cli import main as cli_main
from ocean_research_hub.corpus.repository import CorpusNotFoundError, CorpusRepository
from ocean_research_hub.corpus.workbook import (
    REQUIRED_SHEETS, StagingWorkbook, WorkbookValidationError, normalize_title, parse_claim_value,
    read_workbook,
)
from ocean_research_hub.ingestion.models import PaperWorkflowStatus
from ocean_research_hub.ingestion.repository import SqlitePaperRepository
from ocean_research_hub.schemas.paper_record import PaperRecord

WORKBOOK = Path(__file__).resolve().parents[1] / "import_staging" / "Ocean_Research_Hub_COMPLET.xlsx"


@pytest.fixture(scope="module")
def book() -> StagingWorkbook:
    return read_workbook(WORKBOOK)


@pytest.fixture()
def imported(tmp_path: Path, book: StagingWorkbook) -> CorpusRepository:
    repository = CorpusRepository(tmp_path / "hub.db")
    repository.import_workbook(book)
    return repository


def _edited_workbook(tmp_path: Path, edit) -> Path:
    """Copy the real workbook, apply ``edit(workbook)``, save, return the path."""
    workbook = load_workbook(WORKBOOK)
    edit(workbook)
    target = tmp_path / "edited.xlsx"
    workbook.save(target)
    return target


def _column(sheet, name: str) -> int:
    header = [cell.value for cell in sheet[1]]
    return header.index(name) + 1


def _mutated(book: StagingWorkbook, **changes) -> StagingWorkbook:
    return replace(copy.deepcopy(book), **changes)


# --- the complete corpus is imported ---------------------------------------

def test_the_workbook_has_all_nine_sheets_and_the_declared_counts(book: StagingWorkbook) -> None:
    assert set(REQUIRED_SHEETS) <= set(book.sheet_names)
    assert len(book.sheet_names) == 9
    assert len(book.paper_index) == 115
    assert len(book.paper_records) == 115
    assert len(book.claims) == 144
    assert len(book.field_contract) == 162
    assert len(book.research_gaps) == 8
    assert set(book.documents) == {"README", "AUDIT_PROTOCOL", "METHODS_GUIDE"}


def test_first_import_creates_every_row(tmp_path: Path, book: StagingWorkbook) -> None:
    report = CorpusRepository(tmp_path / "hub.db").import_workbook(book)
    counts = report.counts()
    assert counts["papers"]["created"] == 115
    assert counts["claims"]["created"] == 144
    assert counts["field_contract"]["created"] == 162
    assert counts["research_gaps"]["created"] == 8
    assert counts["documents"]["created"] == 4


def test_every_claim_keeps_its_candidate_status(imported: CorpusRepository) -> None:
    claims, total = imported.list_claims(limit=1000)
    assert total == 144
    assert {claim["scientific_status"] for claim in claims} == {"NOT_VERIFIED"}
    assert {claim["independent_audit"] for claim in claims} == {"PENDING_INDEPENDENT_AUDIT"}
    assert imported.summary()["claim_statuses"] == {"NOT_VERIFIED": 144}


def test_claim_source_columns_are_preserved_exactly(imported: CorpusRepository, book: StagingWorkbook) -> None:
    mapping = {
        "evidence_source_url": ("evidence", "source_url"), "source_doi": ("evidence", "source_doi"),
        "source_edition": ("evidence", "source_edition"), "source_section": ("evidence", "section"),
        "source_locator": ("evidence", "locator"), "verbatim_evidence": ("evidence", "verbatim_evidence"),
        "pdf_page": ("evidence", "pdf_page"), "evidence_review": ("evidence", "evidence_review"),
    }
    for row in book.claims:
        stored = imported.get_claim(row["claim_id"])
        assert stored["value_raw"] == row["value"]
        for column in ("paper_id", "experiment_id", "field_path", "subject_scope", "claim_type",
                       "scientific_status", "independent_audit", "notes"):
            key = "source_paper_id" if column == "paper_id" else column
            assert stored[key] == row[column], (row["claim_id"], column)
        for column, (group, key) in mapping.items():
            assert stored[group][key] == row[column], (row["claim_id"], column)


def test_no_page_number_or_quotation_is_invented(imported: CorpusRepository) -> None:
    """The package performed no PDF alignment: every page and quotation is null."""
    claims, _ = imported.list_claims(limit=1000)
    assert all(claim["evidence"]["pdf_page"] is None for claim in claims)
    assert all(claim["evidence"]["verbatim_evidence"] is None for claim in claims)


def test_empty_fields_are_not_extracted_never_not_reported(imported: CorpusRepository, book: StagingWorkbook) -> None:
    populated_total = 0
    for row in book.paper_index:
        paper = imported.get_paper(row["paper_id"])
        assert len(paper["fields"]) == 162
        for item in paper["fields"]:
            assert item["status"] != "NOT_REPORTED"
            if not item["claims"]:
                assert item["status"] == "NOT_EXTRACTED"
            else:
                assert item["status"] == "NOT_VERIFIED"
        populated_total += sum(len(item["claims"]) for item in paper["fields"])
    assert populated_total == 144


def test_experiments_and_field_contract_are_preserved(imported: CorpusRepository) -> None:
    paper = imported.get_paper("OAI-0006")
    assert paper["experiments"] == ["OAI-0006:FNO3D", "OAI-0006:MAIN_EXPERIMENT", "OAI-0006:RFNO2D"]
    fields = imported.field_contract()
    assert len(fields) == 162
    assert fields[0] == {"field_path": "identity.title", "group": "identity", "field_name": "title"}


def test_list_values_are_parsed_only_from_literal_string_lists() -> None:
    assert parse_claim_value("['nadir altimetry', 'wide-swath altimetry']") == [
        "nadir altimetry", "wide-swath altimetry",
    ]
    assert parse_claim_value("[1, 2]") == "[1, 2]"
    assert parse_claim_value("__import__('os').system('true')") == "__import__('os').system('true')"
    assert parse_claim_value("[broken") == "[broken"


def test_disagreeing_audit_states_are_both_kept_and_reported(imported: CorpusRepository, book: StagingWorkbook) -> None:
    paper = imported.get_paper("OAI-0001")
    assert paper["scientific_audit_status"] == {
        "paper_index": "NOT_AUDITED", "paper_records": "PENDING_INDEPENDENT_AUDIT",
    }
    assert sum("both are preserved" in warning for warning in book.warnings) == 7


def test_research_gaps_are_team_hypotheses_kept_apart_from_claims(imported: CorpusRepository) -> None:
    gaps = imported.research_gaps()
    assert len(gaps) == 8
    assert {gap["status"] for gap in gaps} == {"TEAM_HYPOTHESIS_NOT_VALIDATED"}
    assert {gap["provenance_type"] for gap in gaps} == {"TEAM_NOTE"}
    claims, _ = imported.list_claims(limit=1000)
    assert not any(claim["field_path"].startswith("interpretation.candidate_research_gaps") for claim in claims)


# --- idempotency and deduplication ----------------------------------------

def test_reimport_is_a_no_op(tmp_path: Path, book: StagingWorkbook) -> None:
    repository = CorpusRepository(tmp_path / "hub.db")
    repository.import_workbook(book)
    before = repository.get_paper("OAI-0002")
    report = repository.import_workbook(book)
    counts = report.counts()
    for entity, total in (("papers", 115), ("claims", 144), ("field_contract", 162), ("research_gaps", 8)):
        assert counts[entity] == {"created": 0, "updated": 0, "unchanged": total, "skipped_protected": 0}
    assert repository.summary()["papers"] == 115
    assert repository.get_paper("OAI-0002")["import_run_id"] == before["import_run_id"]
    assert len(repository.import_runs()) == 2


def test_a_changed_unaudited_claim_is_updated(imported: CorpusRepository, book: StagingWorkbook) -> None:
    claims = copy.deepcopy(book.claims)
    claims[0]["notes"] = "revised by a later staging pass"
    report = imported.import_workbook(_mutated(book, claims=claims))
    assert report.claims.updated == 1
    assert imported.get_claim(claims[0]["claim_id"])["notes"] == "revised by a later staging pass"


def test_an_audited_claim_is_never_overwritten(imported: CorpusRepository, book: StagingWorkbook) -> None:
    claim_id = book.claims[0]["claim_id"]
    # Simulate an independent audit recorded after import.
    with sqlite3.connect(imported.database_path) as connection:
        connection.execute(
            "UPDATE staging_corpus_claims SET scientific_status = 'VERIFIED', "
            "independent_audit = 'ATTESTED_BY_AUDITOR' WHERE claim_id = ?", (claim_id,),
        )
    claims = copy.deepcopy(book.claims)
    claims[0]["value"] = "a different value from a later staging pass"
    report = imported.import_workbook(_mutated(book, claims=claims))

    assert report.claims.skipped_protected == 1
    stored = imported.get_claim(claim_id)
    assert stored["value_raw"] == book.claims[0]["value"]
    assert stored["scientific_status"] == "VERIFIED"
    assert any("NOT applied" in warning for warning in report.warnings)


def test_an_audited_paper_is_never_overwritten(imported: CorpusRepository, book: StagingWorkbook) -> None:
    with sqlite3.connect(imported.database_path) as connection:
        row = connection.execute(
            "SELECT record_json FROM staging_corpus_papers WHERE paper_id = 'OAI-0003'"
        ).fetchone()
        record = json.loads(row[0])
        record["scientific_audit_status"] = "INDEPENDENTLY_AUDITED"
        connection.execute(
            "UPDATE staging_corpus_papers SET record_json = ? WHERE paper_id = 'OAI-0003'",
            (json.dumps(record),),
        )
    index = copy.deepcopy(book.paper_index)
    target = next(row for row in index if row["paper_id"] == "OAI-0003")
    target["notes"] = "overwrite attempt"
    report = imported.import_workbook(_mutated(book, paper_index=index))
    assert report.papers.skipped_protected == 1
    assert imported.get_paper("OAI-0003")["paper_index"]["notes"] != "overwrite attempt"


def test_a_new_id_for_an_existing_doi_becomes_an_alias(imported: CorpusRepository, book: StagingWorkbook) -> None:
    index = copy.deepcopy(book.paper_index)
    records = copy.deepcopy(book.paper_records)
    claims = copy.deepcopy(book.claims)
    for row in (*index, *records, *claims):
        if row["paper_id"] == "OAI-0002":
            row["paper_id"] = "EXT-9002"
    report = imported.import_workbook(_mutated(book, paper_index=index, paper_records=records, claims=claims))

    assert report.aliases_created == [{"alias_id": "EXT-9002", "paper_id": "OAI-0002", "matched_on": "doi"}]
    assert imported.summary()["papers"] == 115
    assert imported.get_paper("EXT-9002")["paper_id"] == "OAI-0002"
    # Claims of the aliased id stay attached to the one canonical paper.
    assert {claim["paper_id"] for claim in imported.list_claims(paper_id="OAI-0002", limit=500)[0]} == {"OAI-0002"}
    # The alias resolves in the claim listing too.
    assert imported.list_claims(paper_id="EXT-9002")[1] == imported.list_claims(paper_id="OAI-0002")[1] > 0


def test_title_and_year_deduplicate_only_when_a_doi_is_missing(imported: CorpusRepository, book: StagingWorkbook) -> None:
    no_doi = next(row for row in book.paper_index if not row.get("doi"))
    index = copy.deepcopy(book.paper_index)
    records = copy.deepcopy(book.paper_records)
    for row in (*index, *records):
        if row["paper_id"] == no_doi["paper_id"]:
            row["paper_id"] = "EXT-TITLE"
    report = imported.import_workbook(_mutated(book, paper_index=index, paper_records=records))
    assert {"alias_id": "EXT-TITLE", "paper_id": no_doi["paper_id"], "matched_on": "title_year"} in report.aliases_created
    assert imported.summary()["papers"] == 115

    # Two different DOIs are two works, whatever their titles say.
    same_title = imported.get_paper("OAI-0001")
    connection = imported._ready()  # noqa: SLF001 - identity rule under test
    try:
        resolved = imported._resolve_paper(  # noqa: SLF001
            connection, "EXT-NEW", "10.9999/other-work",
            normalize_title(same_title["title"]), same_title["year"],
        )
    finally:
        connection.close()
    assert resolved == (None, None)


def test_rows_missing_from_a_later_workbook_are_reported_and_kept(imported: CorpusRepository, book: StagingWorkbook) -> None:
    drop = "OAI-0115"
    report = imported.import_workbook(_mutated(
        book,
        paper_index=[row for row in book.paper_index if row["paper_id"] != drop],
        paper_records=[row for row in book.paper_records if row["paper_id"] != drop],
        claims=[row for row in book.claims if row["paper_id"] != drop],
    ))
    assert report.absent_from_source["papers"] == [drop]
    assert len(report.absent_from_source["claims"]) == 4
    assert imported.get_paper(drop)["paper_id"] == drop


def test_a_failed_import_writes_nothing(tmp_path: Path, book: StagingWorkbook, monkeypatch) -> None:
    repository = CorpusRepository(tmp_path / "hub.db")

    def explode(*args, **kwargs):
        raise RuntimeError("simulated failure after papers and claims were written")

    monkeypatch.setattr(repository, "_import_gaps", explode)
    with pytest.raises(RuntimeError):
        repository.import_workbook(book)
    summary = repository.summary()
    assert (summary["papers"], summary["claims"], summary["field_definitions"]) == (0, 0, 0)
    assert repository.import_runs() == []


# --- canonical PaperRecords are linked, never written ----------------------

def test_the_canonical_paper_table_is_linked_but_never_written(tmp_path: Path, book: StagingWorkbook) -> None:
    database = tmp_path / "hub.db"
    papers = SqlitePaperRepository(database)
    papers.initialize()
    stored, _ = papers.create_or_get(
        identity_key="doi:10.5194/gmd-16-2119-2023", doi="10.5194/gmd-16-2119-2023",
        record=PaperRecord(), workflow_status=PaperWorkflowStatus.SCIENTIFIC_AUDIT, warnings=[],
    )
    with sqlite3.connect(database) as connection:
        before = connection.execute("SELECT * FROM papers").fetchall()

    report = CorpusRepository(database).import_workbook(book)

    with sqlite3.connect(database) as connection:
        after = connection.execute("SELECT * FROM papers").fetchall()
    assert after == before
    assert report.linked_paper_records == 1
    assert CorpusRepository(database).get_paper("OAI-0001")["linked_paper_record_id"] == stored.id


# --- validation refuses unsafe or damaged packages -------------------------

def _set_claim(workbook, column: str, value) -> None:
    sheet = workbook["SCIENTIFIC_CLAIMS"]
    sheet.cell(row=2, column=_column(sheet, column)).value = value


@pytest.mark.parametrize(("status", "reason"), [
    ("VERIFIED", "independent attestation"),
    ("NOT_REPORTED", "complete search scope"),
    ("PROBABLY_TRUE", "unknown scientific_status"),
])
def test_a_status_a_staging_package_cannot_carry_is_refused(tmp_path: Path, status: str, reason: str) -> None:
    path = _edited_workbook(tmp_path, lambda wb: _set_claim(wb, "scientific_status", status))
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(path)
    assert any(reason in error for error in caught.value.errors)


def test_a_missing_worksheet_is_refused(tmp_path: Path) -> None:
    path = _edited_workbook(tmp_path, lambda wb: wb.remove(wb["FIELD_CONTRACT"]))
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(path)
    assert caught.value.errors == ["missing worksheet FIELD_CONTRACT"]


def test_a_claim_outside_the_field_contract_is_refused(tmp_path: Path) -> None:
    path = _edited_workbook(tmp_path, lambda wb: _set_claim(wb, "field_path", "architecture.invented_field"))
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(path)
    assert any("outside FIELD_CONTRACT" in error for error in caught.value.errors)


def test_a_truncated_sheet_contradicting_start_here_is_refused(tmp_path: Path) -> None:
    def truncate(workbook) -> None:
        sheet = workbook["RESEARCH_GAPS"]
        sheet.delete_rows(sheet.max_row)
    path = _edited_workbook(tmp_path, truncate)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(path)
    assert any("Research-gap candidates = 8" in error for error in caught.value.errors)


def test_a_doi_shared_by_two_paper_ids_is_refused(tmp_path: Path) -> None:
    def duplicate(workbook) -> None:
        index = workbook["PAPER_INDEX"]
        doi = _column(index, "doi")
        index.cell(row=3, column=doi).value = index.cell(row=2, column=doi).value
        records = workbook["PAPER_RECORDS"]
        record = json.loads(records.cell(row=3, column=2).value)
        record["doi"] = index.cell(row=2, column=doi).value
        records.cell(row=3, column=2).value = json.dumps(record)
    path = _edited_workbook(tmp_path, duplicate)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(path)
    assert any("assigned to several paper_ids" in error for error in caught.value.errors)


def test_a_dangling_claim_reference_is_refused(tmp_path: Path) -> None:
    def dangle(workbook) -> None:
        records = workbook["PAPER_RECORDS"]
        record = json.loads(records.cell(row=2, column=2).value)
        record["field_claim_ids"]["problem"]["geography"] = ["ORI-9999"]
        records.cell(row=2, column=2).value = json.dumps(record)
    path = _edited_workbook(tmp_path, dangle)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(path)
    assert any("ORI-9999" in error for error in caught.value.errors)


# --- CLI ------------------------------------------------------------------

def test_cli_dry_run_writes_nothing(tmp_path: Path, capsys) -> None:
    database = tmp_path / "hub.db"
    assert cli_main([str(WORKBOOK), "--database", str(database), "--dry-run"]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "VALID_NOT_IMPORTED"
    assert not database.exists()


def test_cli_rejects_an_invalid_workbook_without_writing(tmp_path: Path, capsys) -> None:
    path = _edited_workbook(tmp_path, lambda wb: _set_claim(wb, "scientific_status", "VERIFIED"))
    database = tmp_path / "hub.db"
    assert cli_main([str(path), "--database", str(database)]) == 2
    assert json.loads(capsys.readouterr().err)["status"] == "REJECTED"
    assert not database.exists()


def test_cli_imports_and_writes_a_report(tmp_path: Path, capsys) -> None:
    report = tmp_path / "report.json"
    assert cli_main([str(WORKBOOK), "--database", str(tmp_path / "hub.db"), "--report-out", str(report)]) == 0
    assert json.loads(report.read_text(encoding="utf-8"))["counts"]["claims"]["created"] == 144


# --- API and website --------------------------------------------------------

@pytest.fixture()
def client(tmp_path: Path, book: StagingWorkbook):
    database = tmp_path / "hub.db"
    corpus = CorpusRepository(database)
    corpus.import_workbook(book)
    app = create_app(repository=SqlitePaperRepository(database), corpus_repository=corpus)
    with TestClient(app) as test_client:
        yield test_client


def test_api_summary_and_listing(client: TestClient) -> None:
    summary = client.get("/api/corpus").json()
    assert (summary["papers"], summary["claims"], summary["field_definitions"], summary["research_gap_candidates"]) == (115, 144, 162, 8)
    assert "NOT independently verified" in summary["notice"]
    listing = client.get("/api/corpus/papers", params={"limit": 500}).json()
    assert listing["total"] == 115 and len(listing["papers"]) == 115
    assert client.get("/api/corpus/papers", params={"q": "OceanNet"}).json()["total"] == 1


def test_api_paper_detail_claims_fields_and_gaps(client: TestClient) -> None:
    paper = client.get("/api/corpus/papers/OAI-0001").json()
    assert paper["doi"] == "10.5194/gmd-16-2119-2023"
    assert len(paper["fields"]) == 162
    assert paper["populated_field_count"] == 22
    task = next(item for item in paper["fields"] if item["field_path"] == "problem.task")
    assert task["status"] == "NOT_VERIFIED"
    assert task["claims"][0]["evidence"]["source_url"] == "https://gmd.copernicus.org/articles/16/2119/2023/"
    geography = next(item for item in paper["fields"] if item["field_path"] == "problem.geography")
    assert geography["status"] == "NOT_EXTRACTED"

    claims = client.get("/api/corpus/claims", params={"paper_id": "OAI-0001"}).json()
    assert claims["total"] == 22
    claim = client.get("/api/corpus/claims/ORI-0002").json()
    assert claim["value"] == ["nadir altimetry", "wide-swath altimetry"]
    assert client.get("/api/corpus/fields").json()["total"] == 162
    gaps = client.get("/api/corpus/research-gaps").json()
    assert gaps["total"] == 8 and "team hypotheses" in gaps["notice"]
    assert client.get("/api/corpus/import-runs").json()["import_runs"][0]["counts"]["papers"]["created"] == 115


def test_api_unknown_ids_are_404(client: TestClient) -> None:
    assert client.get("/api/corpus/papers/OAI-9999").status_code == 404
    assert client.get("/api/corpus/claims/ORI-9999").status_code == 404
    assert client.get("/corpus/papers/OAI-9999").status_code == 404


def test_website_pages_render_with_the_staging_notice(client: TestClient) -> None:
    index = client.get("/corpus")
    assert index.status_code == 200
    assert "NOT independently verified" in index.text
    assert index.text.count('href="/corpus/papers/') == 115
    assert "Cross-regime robustness" in index.text
    detail = client.get("/corpus/papers/OAI-0001")
    assert detail.status_code == 200
    assert "NOT_EXTRACTED" in detail.text and "ORI-0001" in detail.text
    assert "4DVarNet-SSH" in detail.text


def test_the_existing_api_still_works_beside_the_corpus(client: TestClient) -> None:
    assert client.get("/health").status_code == 200
    assert client.get("/api/landing").status_code == 200


def test_corpus_lookup_raises_for_unknown_paper(imported: CorpusRepository) -> None:
    with pytest.raises(CorpusNotFoundError):
        imported.get_paper("nope")


# --- supplement workbook (EXPANDED_V2) --------------------------------------

SUPPLEMENT_WORKBOOK = WORKBOOK.with_name("Ocean_Research_Intelligence_EXPANDED_V2.xlsx")


@pytest.fixture(scope="module")
def supplement() -> StagingWorkbook:
    return read_workbook(SUPPLEMENT_WORKBOOK)


@pytest.fixture()
def enriched(imported: CorpusRepository, supplement: StagingWorkbook) -> CorpusRepository:
    imported.import_workbook(supplement)
    return imported


def test_the_supplement_is_read_with_its_declared_counts(supplement: StagingWorkbook) -> None:
    assert supplement.kind == "supplement"
    assert (len(supplement.paper_index), len(supplement.claims), len(supplement.research_gaps)) == (115, 258, 8)
    assert (len(supplement.coverage), len(supplement.new_detailed_papers)) == (115, 13)
    assert supplement.paper_records == [] and supplement.field_contract == []


def test_the_supplement_adds_only_new_claims(imported: CorpusRepository, supplement: StagingWorkbook) -> None:
    counts = imported.import_workbook(supplement).counts()
    assert counts["papers"] == {"created": 0, "updated": 0, "unchanged": 115, "skipped_protected": 0}
    assert counts["claims"] == {"created": 114, "updated": 0, "unchanged": 144, "skipped_protected": 0}
    assert counts["research_gaps"]["unchanged"] == 8
    assert counts["aliases_created"] == 0 and counts["absent_from_source"] == {}
    summary = imported.summary()
    assert (summary["papers"], summary["claims"], summary["field_definitions"]) == (115, 258, 162)
    assert summary["claim_statuses"] == {"NOT_VERIFIED": 258}
    assert summary["claim_independent_audit"] == {"PENDING_INDEPENDENT_AUDIT": 144, "PENDING_PDF_AUDIT": 114}
    with sqlite3.connect(imported.database_path) as connection:
        assert connection.execute(
            "SELECT COUNT(*), COUNT(DISTINCT doi_key) FROM staging_corpus_papers WHERE doi_key IS NOT NULL"
        ).fetchone() == (110, 110)


def test_reimporting_the_supplement_is_a_no_op(enriched: CorpusRepository, supplement: StagingWorkbook) -> None:
    counts = enriched.import_workbook(supplement).counts()
    for entity, total in (("papers", 115), ("claims", 258), ("research_gaps", 8), ("coverage", 115), ("documents", 2)):
        assert counts[entity] == {"created": 0, "updated": 0, "unchanged": total, "skipped_protected": 0}
    assert enriched.summary()["claims"] == 258


def test_supplement_claims_keep_their_evidence_and_invent_nothing(enriched: CorpusRepository) -> None:
    claim = enriched.get_claim("WEB2-0001")
    assert claim["scientific_status"] == "NOT_VERIFIED"
    assert claim["independent_audit"] == "PENDING_PDF_AUDIT"
    assert claim["claim_type"] == "AUTHOR_REPORTED_FACT"
    assert claim["evidence"]["source_url"] == "https://arxiv.org/abs/2308.11814"
    assert claim["evidence"]["section"] == "Abstract"
    assert claim["evidence"]["pdf_page"] is None and claim["evidence"]["verbatim_evidence"] is None


def test_supplement_claims_fill_fields_by_path_and_the_rest_stay_not_extracted(enriched: CorpusRepository) -> None:
    paper = enriched.get_paper("OAI-0005")
    task = next(item for item in paper["fields"] if item["field_path"] == "problem.task")
    assert task["status"] == "NOT_VERIFIED" and "WEB2-0001" in task["claim_ids"]
    assert paper["populated_field_count"] == 6
    assert paper["not_extracted_field_count"] == 156
    assert {field["status"] for field in paper["fields"] if not field["claims"]} == {"NOT_EXTRACTED"}
    # The stored record is kept as the base package supplied it.
    assert paper["detail_extraction_status"] == "NOT_EXTRACTED"
    assert paper["coverage"][0]["coverage_level"] == "DETAILED"
    assert paper["coverage"][0]["audit_state"] == "PDF_AUDIT_PENDING"


def test_claims_outside_the_contract_are_kept_but_mapped_to_no_field(enriched: CorpusRepository) -> None:
    paper = enriched.get_paper("OAI-0063")
    assert {claim["field_path"] for claim in paper["uncontracted_claims"]} == {"rl.action", "rl.reward"}
    assert not any("WEB2-0072" in field["claim_ids"] for field in paper["fields"])
    assert "rl.action" not in {field["field_path"] for field in enriched.field_contract()}


def test_a_pending_pdf_audit_claim_is_not_treated_as_attested(enriched: CorpusRepository, supplement: StagingWorkbook) -> None:
    claims = copy.deepcopy(supplement.claims)
    target = next(row for row in claims if row["claim_id"] == "WEB2-0001")
    target["notes"] = "revised"
    report = enriched.import_workbook(_mutated(supplement, claims=claims))
    assert report.claims.updated == 1 and report.claims.skipped_protected == 0


def test_the_supplement_never_overwrites_an_audited_claim(enriched: CorpusRepository, supplement: StagingWorkbook) -> None:
    with sqlite3.connect(enriched.database_path) as connection:
        connection.execute(
            "UPDATE staging_corpus_claims SET scientific_status = 'VERIFIED', "
            "independent_audit = 'AUDITED_BY_SCIENTIFIC_AUDITOR' WHERE claim_id = 'ORI-0001'"
        )
    claims = copy.deepcopy(supplement.claims)
    next(row for row in claims if row["claim_id"] == "ORI-0001")["value"] = "rewritten"
    report = enriched.import_workbook(_mutated(supplement, claims=claims))
    assert report.claims.skipped_protected == 1
    assert enriched.get_claim("ORI-0001")["scientific_status"] == "VERIFIED"
    assert enriched.get_claim("ORI-0001")["value"] != "rewritten"


def test_a_supplement_without_a_base_corpus_is_refused_and_writes_nothing(tmp_path: Path, supplement: StagingWorkbook) -> None:
    repository = CorpusRepository(tmp_path / "hub.db")
    with pytest.raises(WorkbookValidationError) as caught:
        repository.import_workbook(supplement)
    assert any("import the base workbook first" in error for error in caught.value.errors)
    assert repository.summary()["claims"] == 0 and repository.import_runs() == []


def test_a_supplement_cannot_introduce_a_paper(imported: CorpusRepository, supplement: StagingWorkbook) -> None:
    index = copy.deepcopy(supplement.paper_index)
    index[0].update(paper_id="OAI-9999", doi="10.9999/new", title="A paper the base never had")
    claims_before = imported.summary()["claims"]
    with pytest.raises(WorkbookValidationError) as caught:
        imported.import_workbook(_mutated(supplement, paper_index=index))
    assert any("OAI-9999" in error for error in caught.value.errors)
    assert imported.summary()["claims"] == claims_before


def test_a_supplement_whose_coverage_disagrees_with_its_claims_is_refused(tmp_path: Path) -> None:
    workbook = load_workbook(SUPPLEMENT_WORKBOOK)
    sheet = workbook["COVERAGE"]
    sheet.cell(row=2, column=_column(sheet, "claim_count"), value=999)
    target = tmp_path / "edited.xlsx"
    workbook.save(target)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(target)
    assert any("COVERAGE declares 999 claims" in error for error in caught.value.errors)


def test_api_serves_the_enriched_corpus(tmp_path: Path, book: StagingWorkbook, supplement: StagingWorkbook) -> None:
    database = tmp_path / "hub.db"
    corpus = CorpusRepository(database)
    corpus.import_workbook(book)
    corpus.import_workbook(supplement)
    app = create_app(repository=SqlitePaperRepository(database), corpus_repository=corpus)
    with TestClient(app) as client:
        summary = client.get("/api/corpus").json()
        assert (summary["papers"], summary["claims"]) == (115, 258)
        assert client.get("/api/corpus/claims", params={"limit": 1000}).json()["total"] == 258
        assert client.get("/api/corpus/papers/OAI-0063").json()["uncontracted_claims"]
        page = client.get("/corpus/papers/OAI-0063")
        assert "Outside the field contract" in page.text and "rl.action" in page.text


# --- supplement workbook (ENRICHED: B03-B07, field-status markers) -----------

ENRICHED_WORKBOOK = WORKBOOK.with_name("Ocean_Research_Intelligence_ENRICHED.xlsx")


@pytest.fixture(scope="module")
def enriched_book() -> StagingWorkbook:
    return read_workbook(ENRICHED_WORKBOOK)


@pytest.fixture()
def full(enriched: CorpusRepository, enriched_book: StagingWorkbook) -> CorpusRepository:
    enriched.import_workbook(enriched_book)
    return enriched


def test_enriched_is_read_with_its_declared_counts(enriched_book: StagingWorkbook) -> None:
    assert enriched_book.kind == "supplement"
    assert (len(enriched_book.paper_index), len(enriched_book.claims), len(enriched_book.research_gaps)) == (115, 1497, 17)
    assert sum(1 for row in enriched_book.claims if row.get("extraction_status") in {"NOT_EXTRACTED", "NOT_REPORTED"}) == 52
    assert set(enriched_book.supplement_documents) == {"NEW_DETAILED_PAPERS", "SOURCE_ACCESS", "FIELD_COVERAGE"}


def test_enriched_adds_claims_markers_and_gaps_on_top_of_v2(enriched: CorpusRepository, enriched_book: StagingWorkbook) -> None:
    counts = enriched.import_workbook(enriched_book).counts()
    # The 258 earlier claims gain only empty new columns: unchanged, not rewritten.
    assert counts["claims"] == {"created": 1187, "updated": 0, "unchanged": 258, "skipped_protected": 0}
    assert counts["field_markers"]["created"] == 52
    assert counts["research_gaps"] == {"created": 9, "updated": 0, "unchanged": 8, "skipped_protected": 0}
    assert counts["papers"] == {"created": 0, "updated": 67, "unchanged": 48, "skipped_protected": 0}
    assert counts["aliases_created"] == 0 and counts["absent_from_source"] == {}
    summary = enriched.summary()
    assert (summary["papers"], summary["claims"], summary["research_gap_candidates"]) == (115, 1445, 17)
    assert summary["claim_statuses"] == {"NOT_VERIFIED": 1445}
    assert summary["field_status_markers"] == {"NOT_EXTRACTED": 46, "NOT_REPORTED": 6}


def test_reimporting_enriched_is_a_no_op(full: CorpusRepository, enriched_book: StagingWorkbook) -> None:
    counts = full.import_workbook(enriched_book).counts()
    for entity, total in (("papers", 115), ("claims", 1445), ("field_markers", 52), ("research_gaps", 17), ("coverage", 115)):
        assert counts[entity] == {"created": 0, "updated": 0, "unchanged": total, "skipped_protected": 0}


def test_markers_are_never_served_as_values(full: CorpusRepository) -> None:
    fields = {item["field_path"]: item for item in full.get_paper("OAI-0010")["fields"]}
    activation = fields["architecture.activation"]
    assert activation["status"] == "NOT_REPORTED_CANDIDATE" and activation["claims"] == []
    marker = activation["markers"][0]
    assert marker["scientific_status"] == "NOT_VERIFIED"
    assert "appendices" in marker["section_or_scope"]
    # A NOT_REPORTED marker beside an extracted value is a visible conflict.
    assert fields["training.training_time"]["status"] == "NOT_REPORTED_CANDIDATE|NOT_VERIFIED"
    fields = {item["field_path"]: item for item in full.get_paper("OAI-0011")["fields"]}
    assert {item["status"] for item in fields.values() if item["markers"] and not item["claims"]} <= {"NOT_EXTRACTED"}
    assert all(claim["value_raw"] not in ("NOT_EXTRACTED", "NOT_REPORTED") for claim in full.list_claims(limit=2000)[0])


def test_supplied_pdf_pages_and_units_are_kept_as_supplied(full: CorpusRepository) -> None:
    claim = full.get_claim("B03-0001")
    assert claim["evidence"]["pdf_page"] == 6 and claim["batch_id"] == "B03"
    assert claim["evidence"]["verbatim_evidence"] is None
    assert full.summary()["claims_with_supplied_pdf_page"] == 291


def test_shifted_gap_rows_are_realigned_and_traced(full: CorpusRepository) -> None:
    gap = next(item for item in full.research_gaps() if item["candidate_topic"] == "B04-G01")
    assert gap["basis_paper_ids"] == ["OAI-0031"]
    assert gap["testable_question"].startswith("Test tail-sensitive losses")
    assert gap["realigned_from"]["testable_question"] == "OAI-0031"
    assert gap["provenance_type"] == "AI_INTERPRETATION"
    assert next(item for item in full.research_gaps() if item["candidate_topic"] == "Cross-regime robustness")["provenance_type"] == "TEAM_NOTE"


def _edited_enriched(tmp_path: Path, sheet_name: str, row: int, column: str, value) -> Path:
    workbook = load_workbook(ENRICHED_WORKBOOK)
    sheet = workbook[sheet_name]
    sheet.cell(row=row, column=_column(sheet, column), value=value)
    target = tmp_path / "edited.xlsx"
    workbook.save(target)
    return target


def test_a_gap_malformed_in_any_other_way_is_still_refused(tmp_path: Path) -> None:
    workbook = load_workbook(ENRICHED_WORKBOOK)
    sheet = workbook["RESEARCH_GAPS"]
    last = sheet.max_row
    sheet.cell(row=last, column=_column(sheet, "candidate_topic"), value="Not a gap id")
    target = tmp_path / "edited.xlsx"
    workbook.save(target)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(target)
    assert any("cites unknown papers" in error for error in caught.value.errors)


def test_a_marker_carrying_a_value_is_refused(tmp_path: Path, enriched_book: StagingWorkbook) -> None:
    row = 1 + next(position for position, claim in enumerate(enriched_book.claims, start=1) if claim["claim_id"] == "B03-0218")
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(_edited_enriched(tmp_path, "SCIENTIFIC_CLAIMS", row, "value", "GELU"))
    assert any("NOT_REPORTED marker but carries the value 'GELU'" in error for error in caught.value.errors)


def test_an_unknown_extraction_status_is_refused(tmp_path: Path) -> None:
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(_edited_enriched(tmp_path, "SCIENTIFIC_CLAIMS", 300, "extraction_status", "VERIFIED"))
    assert any("unknown extraction_status 'VERIFIED'" in error for error in caught.value.errors)


def test_api_serves_the_full_enriched_corpus(tmp_path: Path, book: StagingWorkbook, supplement: StagingWorkbook, enriched_book: StagingWorkbook) -> None:
    database = tmp_path / "hub.db"
    corpus = CorpusRepository(database)
    for workbook in (book, supplement, enriched_book):
        corpus.import_workbook(workbook)
    app = create_app(repository=SqlitePaperRepository(database), corpus_repository=corpus)
    with TestClient(app) as client:
        summary = client.get("/api/corpus").json()
        assert (summary["papers"], summary["claims"], summary["research_gap_candidates"]) == (115, 1445, 17)
        assert client.get("/api/corpus/claims", params={"limit": 1000, "offset": 1000}).json()["total"] == 1445
        assert "NOT_REPORTED_CANDIDATE" in summary["notice"]
        page = client.get("/corpus/papers/OAI-0010")
        assert page.status_code == 200 and "marker NOT_REPORTED" in page.text and "NOT_REPORTED_CANDIDATE" in page.text


# --- review fixes (QA + Scientific Auditor on PR #27) -----------------------

def _snapshot(repository: CorpusRepository) -> dict:
    return {**{k: v for k, v in repository.summary().items() if k != "last_import"}, "runs": len(repository.import_runs())}


def test_a_workbook_cannot_carry_an_audit_attestation(enriched_book: StagingWorkbook, book: StagingWorkbook, tmp_path: Path) -> None:
    claims = copy.deepcopy(enriched_book.claims)
    claims[-1]["independent_audit"] = "AUDITED_BY_TEAM"
    from ocean_research_hub.corpus import workbook as module
    assert any("AUDITED_BY_TEAM" in error for error in module._claim_row_errors(_mutated(enriched_book, claims=claims)))
    index = copy.deepcopy(book.paper_index)
    index[0]["scientific_audit_status"] = "VERIFIED"
    assert any("not a pending state" in error for error in module._paper_identity_errors(_mutated(book, paper_index=index)))


def test_a_supplement_cannot_change_a_paper_identity(enriched: CorpusRepository, enriched_book: StagingWorkbook) -> None:
    before = _snapshot(enriched)
    for change in ({"doi": "10.9999/fake"}, {"doi": None}, {"title": "Another title"}, {"paper_id": "NEW-1"}):
        index = copy.deepcopy(enriched_book.paper_index)
        target = next(row for row in index if row["paper_id"] == "OAI-0001")
        target.update(change)
        claims = [
            {**row, "paper_id": "NEW-1"} if change.get("paper_id") and row["paper_id"] == "OAI-0001" else row
            for row in enriched_book.claims
        ]
        with pytest.raises(WorkbookValidationError):
            enriched.import_workbook(_mutated(enriched_book, paper_index=index, claims=claims))
    assert _snapshot(enriched) == before
    assert enriched.get_paper("OAI-0001")["doi"] == "10.5194/gmd-16-2119-2023"


def test_an_older_workbook_cannot_revert_newer_content(full: CorpusRepository, book: StagingWorkbook, supplement: StagingWorkbook) -> None:
    before = _snapshot(full)
    for older in (book, supplement):
        with pytest.raises(WorkbookValidationError) as caught:
            full.import_workbook(older)
        assert "older than the stored corpus" in caught.value.errors[0]
    assert _snapshot(full) == before
    assert full.get_paper("OAI-0098")["scientific_extraction_status"] == "PARTIALLY_EXTRACTED"


def test_a_claim_cannot_move_to_another_field_paper_or_kind(full: CorpusRepository, enriched_book: StagingWorkbook) -> None:
    before = _snapshot(full)
    edits = (
        {"claim_id": "ORI-0001", "field_path": "identity.title"},
        {"claim_id": "ORI-0001", "paper_id": "OAI-0002"},
        {"claim_id": "B03-0001", "extraction_status": "NOT_EXTRACTED", "value": "NOT_EXTRACTED"},
    )
    for edit in edits:
        claims = copy.deepcopy(enriched_book.claims)
        next(row for row in claims if row["claim_id"] == edit["claim_id"]).update(edit)
        with pytest.raises(WorkbookValidationError) as caught:
            full.import_workbook(_mutated(enriched_book, claims=claims))
        assert "a new ID is needed" in caught.value.errors[0]
    assert _snapshot(full) == before


def test_a_status_word_is_never_a_claim_value(enriched_book: StagingWorkbook) -> None:
    from ocean_research_hub.corpus import workbook as module
    claims = copy.deepcopy(enriched_book.claims)
    next(row for row in claims if row["claim_id"] == "B03-0001").update(value="NOT_REPORTED")
    assert any("status word 'NOT_REPORTED'" in error for error in module._claim_row_errors(_mutated(enriched_book, claims=claims)))


def test_a_duplicated_column_is_refused(tmp_path: Path) -> None:
    workbook = load_workbook(ENRICHED_WORKBOOK)
    sheet = workbook["SCIENTIFIC_CLAIMS"]
    sheet.cell(row=1, column=sheet.max_column + 1, value="value")
    target = tmp_path / "edited.xlsx"
    workbook.save(target)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(target)
    assert "worksheet SCIENTIFIC_CLAIMS has the column 'value' more than once" in caught.value.errors


def test_a_gap_must_keep_a_question_and_a_not_validated_status(enriched_book: StagingWorkbook) -> None:
    from ocean_research_hub.corpus import workbook as module
    for change in ({"testable_question": None}, {"status": "VALIDATED_BY_TEAM"}):
        gaps = copy.deepcopy(enriched_book.research_gaps)
        gaps[0].update(change)
        assert module._gap_errors(_mutated(enriched_book, research_gaps=gaps))


def test_the_page_shows_units_and_never_links_a_non_http_source(full: CorpusRepository) -> None:
    from ocean_research_hub.corpus.web import _href, render_corpus_paper
    page = render_corpus_paper(full.get_paper("OAI-0010"))
    assert "244.08 <strong>m²/s²</strong>" in page
    assert _href("javascript:alert(1)") == "#" and _href("https://example.org/a") == "https://example.org/a"


def test_dry_run_checks_against_the_stored_corpus_and_writes_nothing(tmp_path: Path, capsys) -> None:
    database = tmp_path / "hub.db"
    assert cli_main([str(WORKBOOK), "--database", str(database)]) == 0
    capsys.readouterr()
    before = CorpusRepository(database).summary()
    assert cli_main([str(SUPPLEMENT_WORKBOOK), "--database", str(database), "--dry-run"]) == 0
    report = json.loads(capsys.readouterr().out)
    assert report["status"] == "VALID_NOT_IMPORTED" and report["counts"]["claims"]["created"] == 114
    assert CorpusRepository(database).summary() == before
    # A supplement without its base is refused by the dry run too.
    assert cli_main([str(SUPPLEMENT_WORKBOOK), "--database", str(tmp_path / "empty.db"), "--dry-run"]) == 2


# --- supplement workbook (ENRICHED_B08) --------------------------------------

B08_WORKBOOK = WORKBOOK.with_name("Ocean_Research_Intelligence_ENRICHED_B08.xlsx")


@pytest.fixture(scope="module")
def b08_book() -> StagingWorkbook:
    return read_workbook(B08_WORKBOOK)


def test_b08_counts_later_markers_as_extracted_rows_and_that_is_reported(b08_book: StagingWorkbook) -> None:
    assert len(b08_book.claims) == 1579
    assert sum(1 for row in b08_book.claims if row.get("extraction_status") in {"NOT_EXTRACTED", "NOT_REPORTED"}) == 59
    assert any("counts 7 marker row(s) as extracted rows" in warning for warning in b08_book.warnings)
    assert set(b08_book.supplement_documents) >= {"CHANGE_LOG", "BATCH_LOG"}


def test_a_count_that_markers_cannot_explain_is_still_refused(tmp_path: Path) -> None:
    workbook = load_workbook(B08_WORKBOOK)
    sheet = workbook["COVERAGE"]
    row = next(r for r in range(2, sheet.max_row + 1) if sheet.cell(row=r, column=_column(sheet, "paper_id")).value == "OAI-0093")
    sheet.cell(row=row, column=_column(sheet, "claim_count"), value=25)
    target = tmp_path / "edited.xlsx"
    workbook.save(target)
    with pytest.raises(WorkbookValidationError) as caught:
        read_workbook(target)
    assert any("COVERAGE declares 25 claims for OAI-0093" in error for error in caught.value.errors)


def test_b08_adds_its_batch_and_is_idempotent(full: CorpusRepository, b08_book: StagingWorkbook, enriched_book: StagingWorkbook) -> None:
    counts = full.import_workbook(b08_book).counts()
    assert counts["claims"] == {"created": 75, "updated": 0, "unchanged": 1445, "skipped_protected": 0}
    assert counts["field_markers"] == {"created": 7, "updated": 0, "unchanged": 52, "skipped_protected": 0}
    assert counts["papers"]["updated"] == 6 and counts["aliases_created"] == 0
    summary = full.summary()
    assert (summary["papers"], summary["claims"]) == (115, 1520)
    assert summary["field_status_markers"] == {"NOT_EXTRACTED": 48, "NOT_REPORTED": 11}
    again = full.import_workbook(b08_book).counts()
    assert again["claims"]["unchanged"] == 1520 and again["papers"]["unchanged"] == 115
    with pytest.raises(WorkbookValidationError):
        full.import_workbook(enriched_book)
